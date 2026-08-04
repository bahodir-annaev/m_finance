"""Import and export routes."""
import html
import os
from datetime import datetime
from flask import Blueprint, request, redirect, send_file
from werkzeug.utils import secure_filename
from auth import require_role
from utils import render_page, t
from models import (
    get_db, get_setting, get_dashboard_data, get_staff_kpi,
    calculate_hourly_rate, calculate_project_cost,
    analyze_excel_for_import, import_excel_data,
)
from import_nizam import import_nizam_file

bp = Blueprint('import_export', __name__)

_SKIP_REASON_LABELS = {
    'no_data': "ustunlar mos keldi, lekin qatorda hech qanday qiymat topilmadi",
    'no_amount': "Приход/расход summasi yo'q (mix import)",
}


def _upload_folder():
    return os.path.join(os.path.dirname(__file__), '..', 'uploads')


@bp.route('/import', methods=['GET', 'POST'])
@require_role('admin')
def import_page():
    result = None
    if request.method == 'POST':
        f = request.files.get('file')
        if f:
            path = os.path.join(_upload_folder(), secure_filename(f.filename))
            f.save(path)
            result = import_nizam_file(path, request.form.get('period', 'all'))

    result_html = ''
    if result:
        unmatched = result.get('unmatched_names', [])
        warn_html = ''
        if unmatched:
            warn_html = f'''<div class="alert" style="background:#fff3cd;border-left:4px solid #ffc107;padding:12px;margin-top:8px;">
              <b>&#x26A0;&#xFE0F; Mos kelmaydigan xodimlar ({len(unmatched)}):</b> {", ".join(unmatched)}<br>
              Bu ismlar NIZAM_MAP da topilmadi.
            </div>'''
        result_html = f'''<div class="alert alert-success">
          {t("imp_success", result["projects"], result["hour_entries"], result["skipped"], result["employees_mapped"])}
        </div>{warn_html}'''

    return render_page('import', 'import.html', result_html=result_html)


@bp.route('/api/import-excel/analyze', methods=['POST'])
@require_role('admin')
def api_import_excel_analyze():
    f = request.files.get('file')
    target = request.form.get('target', 'auto')
    if not f:
        return redirect('/import')

    path = os.path.join(_upload_folder(), secure_filename(f.filename))
    f.save(path)

    try:
        sheets = analyze_excel_for_import(path)
    except Exception as e:
        return f'<html><body style="font-family:Arial;padding:40px"><h2 style="color:red">Xatolik</h2><p>{e}</p><a href="/import">Orqaga</a></body></html>'

    total_imported = 0
    total_skipped = 0
    import_log = []
    unresolved_projects = 0
    unresolved_counterparties = 0
    skip_entries = []

    for sheet in sheets:
        tbl = (('external_transactions' if sheet['type'] == 'external'
                else ('internal_transactions' if sheet['type'] == 'internal'
                else ('mixed_transactions' if sheet['type'] == 'mix' else None)))
               if target == 'auto' else target)
        if not tbl:
            import_log.append(f"Sheet '{sheet['name']}': tur aniqlanmadi, o'tkazildi")
            continue
        col_map = {str(c['index']): c['matched'] for c in sheet['columns'] if c['matched']}
        if not col_map:
            import_log.append(f"Sheet '{sheet['name']}': ustunlar mos kelmadi")
            continue
        res = import_excel_data(path, sheet['name'], tbl, col_map, sheet['header_idx'])
        total_imported += res['imported']
        total_skipped += res['skipped']
        unresolved_projects = res.get('unresolved_projects', 0)
        unresolved_counterparties = res.get('unresolved_counterparties', 0)
        for entry in res.get('skip_log', []):
            skip_entries.append({**entry, 'sheet': sheet['name']})
        import_log.append(f"Sheet '{sheet['name']}' → {tbl}: {res['imported']} qator, {res['skipped']} o'tkazildi")

    log_html = ''.join(f'<li>{l}</li>' for l in import_log)
    warn_html = ''
    if unresolved_projects or unresolved_counterparties:
        warn_html = f'''<div style="background:#fff3cd;border-left:4px solid #ffc107;padding:12px;margin-top:12px">
            <b>&#x26A0;&#xFE0F; Moslashtirish kerak:</b>
            {unresolved_projects} loyiha nomi va {unresolved_counterparties} kontragent nomi
            mavjud ro'yxatlarga avtomatik moslashtirilmadi (project_id/counterparty_id bo'sh qoldi).<br>
            <a href="/reference/reconcile">Moslashtirish sahifasiga o'tish →</a>
        </div>'''

    skip_html = ''
    if skip_entries:
        items = ''.join(
            f"<li><b>{html.escape(e['sheet'])}</b> — qator {e['row']}: "
            f"{_SKIP_REASON_LABELS.get(e['reason'], e['reason'])}"
            f"<br><code style=\"font-size:11px;color:#666\">{html.escape(e['raw']) or '(boʻsh)'}</code></li>"
            for e in skip_entries
        )
        skip_html = f'''<div style="background:#f8f9fa;border-left:4px solid #6c757d;padding:12px;margin-top:12px">
            <b>O'tkazib yuborilgan qatorlar sababi ({len(skip_entries)}{'+' if len(skip_entries) >= 200 else ''}):</b>
            <span style="font-size:12px;color:#666"> (bo'sh qatorlar bu ro'yxatda ko'rsatilmaydi)</span>
            <ul style="line-height:1.6;max-height:280px;overflow:auto;margin-top:8px;padding-left:20px">{items}</ul>
        </div>'''

    return f'''<html><body style="font-family:Arial;padding:40px;max-width:800px">
        <h2 style="color:#00382F">Import yakunlandi</h2>
        <p>Jami import: <b>{total_imported}</b> | O'tkazildi: {total_skipped}</p>
        <ul style="line-height:1.8">{log_html}</ul>
        {warn_html}
        {skip_html}
        <a href="/import">Orqaga qaytish</a></body></html>'''


@bp.route('/export')
@require_role('manager', 'admin')
def export_page():
    return render_page('export', 'export.html')


@bp.route('/api/export/xlsx')
@require_role('manager', 'admin')
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    brown_fill = PatternFill('solid', fgColor='6B5740')
    header_font = Font(bold=True, color='FFFFFF', name='Arial', size=11)

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font = header_font
            cell.fill = brown_fill
            cell.alignment = Alignment(horizontal='center')

    ws = wb.active
    ws.title = '09-Dashboard'
    data = get_dashboard_data()
    ws.append(['MIZAN Finance v4.0 — Dashboard'])
    ws.append([f'Sana: {datetime.now().strftime("%Y-%m-%d %H:%M")}'])
    ws.append([])
    ws.append(['Parametr', 'Qiymat'])
    style_header(ws, 4)
    for k, v in [
        ('Jami proektlar', data['total_projects']),
        ('Jami billable soatlar', data['total_hours']),
        ('Ishchi xodimlar', data['production_staff']),
        ('Admin xodimlar', data['admin_staff']),
        ("O'rtacha Cost Rate (UZS)", data['avg_cost_rate']),
        ("O'rtacha Billing Rate (UZS)", data['avg_billing_rate']),
        ('Utilization %', f"{data['utilization_pct']:.1f}%"),
        ('Jami MIZAN xarajat', data['total_mizan_cost']),
        ('Oylik overhead', data['monthly_overhead']),
    ]:
        ws.append([k, v])
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 20

    ws2 = wb.create_sheet('03-Maosh')
    ws2.append(['Ism', 'Lavozim', "Bo'lim", 'Turi', 'Maosh', 'Premiya', 'Soliq', 'JSSM', 'JAMI'])
    style_header(ws2)
    conn = get_db()
    staff = conn.execute('''SELECT s.*, sh.base_salary, sh.premium FROM staff s
        LEFT JOIN salary_history sh ON s.id = sh.staff_id AND sh.end_date IS NULL
        ORDER BY s.staff_type, s.name''').fetchall()
    conn.close()
    tax_r = get_setting('tax_rate')
    soc_r = get_setting('social_rate')
    for s in staff:
        b = s['base_salary'] or 0; p = s['premium'] or 0; g = b + p
        ws2.append([s['name'], s['role'], s['department'], s['staff_type'], b, p, g * tax_r, g * soc_r, g + g * tax_r + g * soc_r])

    ws3 = wb.create_sheet('07-Soat narxi')
    ws3.append(['Ism', 'Lavozim', "Bo'lim", 'Maosh+Prem', 'Soliq+JSSM', 'Admin (prop)', 'Texn+Lits',
                'Gen.equip', 'Overhead (prop)', 'JAMI oylik', 'Cost Rate', 'Billing Rate', 'Billing USD'])
    style_header(ws3)
    conn = get_db()
    prod_staff = conn.execute("SELECT id, name, role, department FROM staff WHERE staff_type='production' AND is_active=1 ORDER BY name").fetchall()
    conn.close()
    for s in prod_staff:
        r = calculate_hourly_rate(s['id'])
        if not isinstance(r, dict):
            continue
        ws3.append([s['name'], s['role'], s['department'],
                    r['base_salary'] + r['premium'], r['tax'] + r['social'], r['admin_share'],
                    r['personal_eq'] + r['personal_lic'], r['general_eq'], r['overhead_share'],
                    r['total_monthly'], r['cost_rate'], r['billing_rate'], r['billing_usd']])

    ws4 = wb.create_sheet('08-Byudjet')
    ws4.append(['Proekt', 'Soatlar', 'MIZAN cost', 'Billing value', 'Outsourcing', 'Tushum', 'Foyda', 'Margin %'])
    style_header(ws4)
    conn = get_db()
    projs = conn.execute("SELECT id FROM projects WHERE is_billable=1").fetchall()
    conn.close()
    for p in projs:
        pc = calculate_project_cost(p['id'])
        if not pc or pc['total_hours'] == 0:
            continue
        ws4.append([pc['project']['name'], pc['total_hours'], pc['mizan_cost'], pc['mizan_billing'],
                    pc['outsourcing'], pc['income'], pc['profit'], pc['margin']])

    ws5 = wb.create_sheet('10-KPI')
    ws5.append(['Ism', 'Lavozim', "Bo'lim", 'Proektlar', 'Soatlar', 'Cost Rate', 'Billing Rate',
                'Cost Value', 'Revenue Value', 'Revenue USD', 'Utilization %'])
    style_header(ws5)
    for k in get_staff_kpi():
        ws5.append([k['name'], k['role'], k['department'], k['projects'], k['hours'],
                    k['cost_rate'], k['billing_rate'], k['cost_value'], k['revenue_value'],
                    k['revenue_usd'], k['utilization'] * 100])

    for ws_obj in wb.worksheets:
        for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M']:
            if ws_obj.column_dimensions[col_letter].width < 14:
                ws_obj.column_dimensions[col_letter].width = 14

    SHEET_MAP = {'staff': '03-Maosh', 'hourly': '07-Soat narxi',
                 'budget': '08-Byudjet', 'kpi': '10-KPI', 'dashboard': '09-Dashboard'}
    requested_sheet = request.args.get('sheet')
    if requested_sheet in SHEET_MAP:
        for i, ws_obj in enumerate(wb.worksheets):
            if ws_obj.title == SHEET_MAP[requested_sheet]:
                wb.active = i
                break

    path = os.path.join(_upload_folder(), 'mizan_export_v4.xlsx')
    wb.save(path)
    return send_file(path, as_attachment=True, download_name='mizan_finance_v4_export.xlsx')
