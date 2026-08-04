"""NIZAM CRM dan table.xlsx import qilish + Texnika/Litsenziya seed data (v4.0)
XATO #3: Staff ID tizimi + yaxshilangan mapping + unmatched names tracking
"""
import openpyxl
import sys
import os
sys.path.insert(0, os.path.dirname(__file__))
from database import get_db, init_db

# NIZAM → V8 mapping
NIZAM_MAP = {
    "ABBOS ABDULLAYEV": ("Abbos", "BIM Architect", "Arxitektura", "production"),
    "ABDULLA ABDULLAYEV": ("Abdulla", "INT 3D Viz", "Vizualizatsiya", "production"),
    "ABDURASHID ABDUG'OFUROV": ("Abdurashid", "BIM Architect", "BIM", "production"),
    "ADIZ SAIDOV": ("Adiz", "Graphic Designer", "Grafika", "production"),
    "Aziz Omonov": ("Aziz O", "PM", "Boshqaruv", "admin"),
    "AZIZ MUXAMEDOV": ("Aziz M", "Interior Owner", "Interyer", "production"),
    "BEHZOD NIYAZOV": ("Behzod", "Lead Architect", "Arxitektura", "production"),
    "DAVRON DJURAYEV": ("Davron", "System/Digital Director", "IT", "admin"),
    "FARHOD INAGAMOV": ("Farxod", "Architecture Owner", "Arxitektura", "production"),
    "IBROXIM ISLOMOV": ("Ibrohim", "BIM Architect", "BIM", "production"),
    "ISKANDAR XUDOYBERDIYEV": ("Iskandar", "BIM Architect", "BIM", "production"),
    "ISLOM DJURAYEV": ("Islom", "Head of Viz/VR", "Vizualizatsiya", "production"),
    "OLCHINBEK OLIMOV": ("Olchin", "INT BIM", "BIM", "production"),
    "RAMZIDDIN MUXUTDINOV": ("Ramziddin", "3D Visualizator", "Vizualizatsiya", "production"),
    "SHAHZOD RAHMATOV": ("Shaxzod", "Urban/Masterplan", "Arxitektura", "production"),
    "SHAXBOZ UMATALIYEV": ("Shaxboz", "3D Modeler", "Vizualizatsiya", "production"),
    "UMAR SHARIPOV": ("Umar", "HR/Office Manager", "HR", "admin"),
    "UMID RAXMATOV": ("Umid", "Urban Architect", "Arxitektura", "production"),
    "XASAN G'ANIXO'JAYEV": ("Hasan", "INT BIM", "BIM", "production"),
    "ZAFARJON RAXMATOV": ("Zafar", "BIM Head", "BIM", "production"),
    "ZIYOVIDDIN FAXRIDDINOV": ("Ziyovuddin", "3D Visualizator", "Vizualizatsiya", "production"),
    "ZOIR AHMADALIYEV": ("Zoir", "Lead Engineer/GIP", "Muhandislik", "production"),
}

NON_BILLABLE = [
    "office", "hr", "ijtimoiy tarmoq", "задачи компании", "nizam ui design",
    "web site mizan", "3d model trade", "tashkiliy ishlar", "website portfolio",
    "unreal engine", "architecture diagrams", "sanjar test", "123456",
    "daho office", "entrance",
]

DEFAULT_SALARIES = {
    "Farxod": (18000000, 3000000), "Aziz M": (16000000, 2500000),
    "Shaxzod": (14000000, 2000000), "Behzod": (22000000, 5000000),
    "Zoir": (18000000, 3000000), "Umid": (10000000, 0),
    "Zafar": (16000000, 3000000), "Abdurashid": (8000000, 0),
    "Ibrohim": (8000000, 0), "Iskandar": (8000000, 0),
    "Hasan": (8000000, 0), "Olchin": (8000000, 0),
    "Abbos": (6000000, 0), "Islom": (14000000, 2000000),
    "Shaxboz": (10000000, 0), "Ramziddin": (8000000, 0),
    "Ziyovuddin": (8000000, 0), "Abdulla": (7000000, 0),
    "Adiz": (8000000, 0), "Qodir": (22000000, 5000000),
    "Behzod G": (12000000, 0), "Murod": (10000000, 0),
    "Bahodir": (10000000, 0), "Jahongir": (30000000, 10000000),
    "Aziz O": (16000000, 3000000), "Davron": (14000000, 2000000),
    "Sardor": (10000000, 0), "Ziyodulla": (14000000, 3000000),
    "Nurmuhammad": (8000000, 0), "Umar": (8000000, 0),
    "Maftuna": (5000000, 0), "Ruslana": (10000000, 0),
    "Nazokat": (4000000, 0),
}

# Shaxsiy texnika — har xodimga biriktirilgan kompyuter (AUDIT FIX #1)
# (xodim_nomi, jihoz_nomi, narxi UZS, muddat oy)
PERSONAL_EQUIPMENT = [
    ("Farxod", "Dell Precision 7780", 25000000, 36),
    ("Farxod", "Dell U2723QE Monitor", 8000000, 48),
    ("Aziz M", "MacBook Pro 16 M2", 28000000, 36),
    ("Aziz M", "LG 27UK850 Monitor", 6000000, 48),
    ("Shaxzod", "Dell Precision 5680", 22000000, 36),
    ("Behzod", "Dell Precision 7780", 25000000, 36),
    ("Behzod", "Dell U2723QE Monitor", 8000000, 48),
    ("Zoir", "Dell Precision 5680", 22000000, 36),
    ("Umid", "Dell Inspiron 16", 12000000, 36),
    ("Zafar", "HP ZBook Fury 16", 28000000, 36),
    ("Zafar", "Dell U2723QE Monitor", 8000000, 48),
    ("Abdurashid", "Dell Inspiron 16", 12000000, 36),
    ("Ibrohim", "Dell Inspiron 16", 12000000, 36),
    ("Iskandar", "Dell Inspiron 16", 12000000, 36),
    ("Hasan", "Dell Inspiron 16", 12000000, 36),
    ("Olchin", "Dell Inspiron 16", 12000000, 36),
    ("Abbos", "Dell Inspiron 14", 8000000, 36),
    ("Islom", "HP ZBook Fury 16", 28000000, 36),
    ("Islom", "Dell U2723QE Monitor", 8000000, 48),
    ("Shaxboz", "HP ZBook Fury 16", 28000000, 36),
    ("Ramziddin", "Dell Inspiron 16", 12000000, 36),
    ("Ziyovuddin", "Dell Inspiron 16", 12000000, 36),
    ("Abdulla", "Dell Inspiron 14", 8000000, 36),
    ("Adiz", "MacBook Pro 14 M2", 22000000, 36),
    ("Qodir", "Dell Precision 5680", 22000000, 36),
    ("Behzod G", "Dell Inspiron 16", 12000000, 36),
    ("Murod", "Dell Inspiron 14", 8000000, 36),
    ("Bahodir", "Dell Inspiron 14", 8000000, 36),
]

# Shaxsiy litsenziyalar — xodimga biriktirilgan (AUDIT FIX #7)
# (xodim_nomi, litsenziya_nomi, yillik narxi UZS, turi)
PERSONAL_LICENSES = [
    ("Farxod", "Autodesk Revit", 12000000, "named"),
    ("Farxod", "AutoCAD", 8000000, "named"),
    ("Behzod", "Autodesk Revit", 12000000, "named"),
    ("Behzod", "AutoCAD", 8000000, "named"),
    ("Shaxzod", "Autodesk Revit", 12000000, "named"),
    ("Shaxzod", "AutoCAD", 8000000, "named"),
    ("Umid", "AutoCAD", 8000000, "named"),
    ("Zoir", "Autodesk Revit", 12000000, "named"),
    ("Zoir", "AutoCAD", 8000000, "named"),
    ("Zafar", "Autodesk Revit", 12000000, "named"),
    ("Zafar", "Navisworks", 8000000, "named"),
    ("Abdurashid", "Autodesk Revit", 12000000, "named"),
    ("Ibrohim", "Autodesk Revit", 12000000, "named"),
    ("Iskandar", "Autodesk Revit", 12000000, "named"),
    ("Hasan", "Autodesk Revit", 12000000, "named"),
    ("Olchin", "Autodesk Revit", 12000000, "named"),
    ("Abbos", "Autodesk Revit", 12000000, "named"),
    ("Islom", "3ds Max", 10000000, "named"),
    ("Islom", "V-Ray", 6000000, "named"),
    ("Shaxboz", "3ds Max", 10000000, "named"),
    ("Shaxboz", "V-Ray", 6000000, "named"),
    ("Ramziddin", "3ds Max", 10000000, "named"),
    ("Ramziddin", "V-Ray", 6000000, "named"),
    ("Ziyovuddin", "3ds Max", 10000000, "named"),
    ("Ziyovuddin", "V-Ray", 6000000, "named"),
    ("Abdulla", "3ds Max", 10000000, "named"),
    ("Abdulla", "V-Ray", 6000000, "named"),
    ("Aziz M", "3ds Max", 10000000, "named"),
    ("Aziz M", "SketchUp Pro", 4000000, "named"),
    ("Adiz", "Adobe Creative Cloud", 8000000, "named"),
    ("Qodir", "Autodesk Revit", 12000000, "named"),
    ("Qodir", "AutoCAD", 8000000, "named"),
    ("Behzod G", "ETABS", 15000000, "named"),
    ("Murod", "Autodesk Revit MEP", 12000000, "named"),
    ("Bahodir", "AutoCAD Electrical", 10000000, "named"),
]

# Umumiy texnika (printer, server, etc.)
GENERAL_EQUIPMENT = [
    ("HP LaserJet Pro M428", 2, 8000000, 48),
    ("HP Color LaserJet Pro", 1, 12000000, 48),
    ("Dell PowerEdge T40 Server", 1, 18000000, 60),
    ("APC Smart-UPS 1500VA", 3, 5000000, 60),
    ("Synology NAS DS920+", 1, 12000000, 60),
    ("WiFi Router Mikrotik", 2, 3000000, 48),
    ("Ofis mebeli (umumiy)", 1, 30000000, 60),
    ("Plotter HP DesignJet T630", 1, 15000000, 60),
]


def parse_hhmm(val):
    if not val or val == 0:
        return 0.0
    s = str(val).strip()
    if ':' not in s:
        return 0.0
    parts = s.split(':')
    return round(int(parts[0]) + int(parts[1]) / 60, 1)


def ensure_staff_exists(conn):
    """Barcha xodimlarni bazaga qo'shish"""
    # NIZAM dagi xodimlar
    for nizam_name, (v8_name, role, dept, stype) in NIZAM_MAP.items():
        conn.execute('''INSERT OR IGNORE INTO staff (name, nizam_name, role, department, staff_type)
                       VALUES (?, ?, ?, ?, ?)''', (v8_name, nizam_name, role, dept, stype))

    # NIZAM da yo'q xodimlar
    extra = [
        ("Qodir", "Lead Architect", "Arxitektura", "production"),
        ("Behzod G", "Static Engineer", "Muhandislik", "production"),
        ("Murod", "Mechanic Engineer", "Muhandislik", "production"),
        ("Bahodir", "Electric Engineer", "Muhandislik", "production"),
        ("Jahongir", "CEO", "Boshqaruv", "admin"),
        ("Sardor", "Export & BD Lead", "Marketing", "admin"),
        ("Ziyodulla", "Finance Director", "Moliya", "admin"),
        ("Nurmuhammad", "Accountant", "Moliya", "admin"),
        ("Maftuna", "Admin Support", "HR", "admin"),
        ("Ruslana", "BOQ/Smeta", "Moliya", "admin"),
        ("Nazokat", "Cleaner", "HR", "admin"),
    ]
    for name, role, dept, stype in extra:
        conn.execute('''INSERT OR IGNORE INTO staff (name, role, department, staff_type)
                       VALUES (?, ?, ?, ?)''', (name, role, dept, stype))
    conn.commit()

    # Maoshlar
    staff = conn.execute("SELECT id, name FROM staff").fetchall()
    for s in staff:
        existing = conn.execute("SELECT id FROM salary_history WHERE staff_id=? AND end_date IS NULL", (s['id'],)).fetchone()
        if not existing and s['name'] in DEFAULT_SALARIES:
            sal, prem = DEFAULT_SALARIES[s['name']]
            conn.execute("INSERT INTO salary_history (staff_id, base_salary, premium, start_date) VALUES (?,?,?,?)",
                        (s['id'], sal, prem, '2025-06-01'))
    conn.commit()


def seed_equipment_and_licenses(conn):
    """Texnika va litsenziya ma'lumotlarini bazaga kiritish (AUDIT FIX #1, #7)"""
    # Shaxsiy texnika
    existing_eq = conn.execute("SELECT COUNT(*) FROM personal_equipment").fetchone()[0]
    if existing_eq == 0:
        for staff_name, eq_name, price, lifespan in PERSONAL_EQUIPMENT:
            staff = conn.execute("SELECT id FROM staff WHERE name=?", (staff_name,)).fetchone()
            if staff:
                conn.execute(
                    "INSERT INTO personal_equipment (name, staff_id, price, lifespan_months) VALUES (?,?,?,?)",
                    (eq_name, staff['id'], price, lifespan))

    # Umumiy texnika
    existing_gen = conn.execute("SELECT COUNT(*) FROM general_equipment").fetchone()[0]
    if existing_gen == 0:
        for eq_name, qty, price, lifespan in GENERAL_EQUIPMENT:
            conn.execute(
                "INSERT INTO general_equipment (name, quantity, price, lifespan_months) VALUES (?,?,?,?)",
                (eq_name, qty, price, lifespan))

    # Shaxsiy litsenziyalar
    existing_lic = conn.execute("SELECT COUNT(*) FROM personal_licenses").fetchone()[0]
    if existing_lic == 0:
        for staff_name, lic_name, annual_cost, lic_type in PERSONAL_LICENSES:
            staff = conn.execute("SELECT id FROM staff WHERE name=?", (staff_name,)).fetchone()
            if staff:
                conn.execute(
                    "INSERT INTO personal_licenses (name, staff_id, annual_cost, license_type) VALUES (?,?,?,?)",
                    (lic_name, staff['id'], annual_cost, lic_type))

    conn.commit()


def _fuzzy_match_nizam_name(nizam_name):
    """XATO #3: Yaxshilangan ism moslashtirish — case-insensitive va bo'sh joy normalizatsiya"""
    # 1. Aniq moslik
    if nizam_name in NIZAM_MAP:
        return NIZAM_MAP[nizam_name]
    # 2. Case-insensitive moslik
    upper = nizam_name.upper().strip()
    for key, val in NIZAM_MAP.items():
        if key.upper().strip() == upper:
            return val
    # 3. Partial moslik (familiya va ism)
    parts = upper.split()
    if len(parts) >= 2:
        for key, val in NIZAM_MAP.items():
            key_parts = key.upper().split()
            # Familiya va ismning birinchi 3 harfi mos kelsa
            if len(key_parts) >= 2 and key_parts[0][:3] == parts[0][:3] and key_parts[1][:3] == parts[1][:3]:
                return val
    return None


def import_nizam_file(filepath, period='all'):
    """NIZAM CRM table.xlsx import
    XATO #3: Yaxshilangan mapping + unmatched_names ro'yxati

    Does NOT auto-create staff — employees that don't match an existing
    staff.nizam_name / NIZAM_MAP name come back in unmatched_names for the
    admin to resolve manually (add/rename the staff row, then re-import)."""
    init_db()
    conn = get_db()

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    # Parse employee columns
    employees = {}
    for col in range(4, ws.max_column + 1):
        v = ws.cell(2, col).value
        if v:
            employees[col] = v.replace('\n', ' ')

    # XATO #3: Yaxshilangan staff name→id map
    staff_map = {}
    unmatched_names = []
    for nizam_name in employees.values():
        # 1. Avval nizam_name ustuni bo'yicha qidirish (aniq DB moslik)
        row = conn.execute("SELECT id FROM staff WHERE nizam_name=?", (nizam_name,)).fetchone()
        if row:
            staff_map[nizam_name] = row['id']
            continue
        # 2. NIZAM_MAP orqali (yaxshilangan fuzzy matching)
        mapped = _fuzzy_match_nizam_name(nizam_name)
        if mapped:
            v8_name = mapped[0]
            row = conn.execute("SELECT id FROM staff WHERE name=?", (v8_name,)).fetchone()
            if row:
                staff_map[nizam_name] = row['id']
                # DB ga nizam_name ni saqlash (keyingi safar tezroq topiladi)
                conn.execute("UPDATE staff SET nizam_name=? WHERE id=?", (nizam_name, row['id']))
                continue
        # 3. Mos kelmadi — logga yozish
        unmatched_names.append(nizam_name)

    imported_projects = 0
    imported_hours = 0
    skipped = 0

    for row in range(4, ws.max_row + 1):
        name = ws.cell(row, 2).value
        if not name:
            continue
        name_str = str(name)

        # Non-billable filter
        is_non = any(kw in name_str.lower() for kw in NON_BILLABLE)
        total_hrs = parse_hhmm(ws.cell(row, 3).value)

        if is_non or total_hrs == 0:
            skipped += 1
            continue

        # Create/get project
        conn.execute("INSERT OR IGNORE INTO projects (name) VALUES (?)", (name_str,))
        proj = conn.execute("SELECT id FROM projects WHERE name=?", (name_str,)).fetchone()
        if not proj:
            continue
        proj_id = proj['id']
        imported_projects += 1

        # Import hours per employee
        for col, emp_name in employees.items():
            hrs = parse_hhmm(ws.cell(row, col).value)
            if hrs > 0 and emp_name in staff_map:
                # Upsert, NOT INSERT OR REPLACE: REPLACE deletes-and-reinserts the
                # row, silently wiping phase_id and the applied_* rate snapshots
                # on every re-import.
                conn.execute('''INSERT INTO project_hours (project_id, staff_id, hours, period)
                               VALUES (?, ?, ?, ?)
                               ON CONFLICT(project_id, staff_id, period)
                               DO UPDATE SET hours=excluded.hours''',
                            (proj_id, staff_map[emp_name], hrs, period))
                imported_hours += 1

    conn.commit()
    conn.close()

    return {
        'projects': imported_projects,
        'hour_entries': imported_hours,
        'skipped': skipped,
        'employees_mapped': len(staff_map),
        'unmatched_names': unmatched_names,  # XATO #3: Mos kelmaganlar ro'yxati
    }
