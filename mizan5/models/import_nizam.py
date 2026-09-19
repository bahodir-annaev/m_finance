"""NIZAM CRM timesheet import — the monthly `table.xlsx` export.

Sheet layout (as NIZAM exports it): row 2 carries employee names from column
D onward, each data row from row 4 is one project — name in column B, total
hours in column C, then one HH:MM cell per employee.

Two rules carried over from v4, both deliberate:

* Employees are never created by an import. A name that matches neither
  `staff.nizam_name`, a staff name, nor the legacy NIZAM_MAP comes back in
  `unmatched_names` for someone to fix on /staff and re-import. Creating a
  staff row from a spreadsheet header would give it no salary and no rate.
* Non-billable rows (office, HR, internal tooling…) are skipped, and every
  skipped row is reported with its reason — never dropped silently.

Projects ARE created when unknown, exactly as v4 did, because a new client
job first appears in the timesheet. Existing ones are matched through
project_aliases and a normalised name, so a renamed project does not fork.
"""
import os

from .base import get_db, now_ts, resolve_entity_id, _write_audit

# Legacy NIZAM → staff mapping kept from v4 for firms migrating with the same
# people. `staff.nizam_name` is always consulted first; this is the fallback.
NIZAM_MAP = {
    "ABBOS ABDULLAYEV": "Abbos",
    "ABDULLA ABDULLAYEV": "Abdulla",
    "ABDURASHID ABDUG'OFUROV": "Abdurashid",
    "ADIZ SAIDOV": "Adiz",
    "Aziz Omonov": "Aziz O",
    "AZIZ MUXAMEDOV": "Aziz M",
    "BEHZOD NIYAZOV": "Behzod",
    "DAVRON DJURAYEV": "Davron",
    "FARHOD INAGAMOV": "Farxod",
    "IBROXIM ISLOMOV": "Ibrohim",
    "ISKANDAR XUDOYBERDIYEV": "Iskandar",
    "ISLOM DJURAYEV": "Islom",
    "OLCHINBEK OLIMOV": "Olchin",
    "RAMZIDDIN MUXUTDINOV": "Ramziddin",
    "SHAHZOD RAHMATOV": "Shaxzod",
    "SHAXBOZ UMATALIYEV": "Shaxboz",
    "UMAR SHARIPOV": "Umar",
    "UMID RAXMATOV": "Umid",
    "XASAN G'ANIXO'JAYEV": "Hasan",
    "ZAFARJON RAXMATOV": "Zafar",
    "ZIYOVIDDIN FAXRIDDINOV": "Ziyovuddin",
    "ZOIR AHMADALIYEV": "Zoir",
}

NON_BILLABLE = [
    "office", "hr", "ijtimoiy tarmoq", "задачи компании", "nizam ui design",
    "web site mizan", "3d model trade", "tashkiliy ishlar", "website portfolio",
    "unreal engine", "architecture diagrams", "sanjar test", "123456",
    "daho office", "entrance",
]

SKIP_REASONS = ('non_billable', 'no_hours', 'no_name')


def parse_hhmm(val):
    """'12:30' → 12.5; a bare number is taken as hours; blanks are zero."""
    if val is None or val == '':
        return 0.0
    if isinstance(val, (int, float)):
        return round(float(val), 1)
    s = str(val).strip()
    if ':' in s:
        parts = s.split(':')
        try:
            return round(int(parts[0]) + int(parts[1]) / 60, 1)
        except (ValueError, IndexError):
            return 0.0
    try:
        return round(float(s.replace(',', '.')), 1)
    except ValueError:
        return 0.0


def _fuzzy_map(nizam_name):
    """Legacy map lookup: exact, then case-insensitive, then first-3-letters of
    surname and given name (NIZAM exports are not consistent about case)."""
    if nizam_name in NIZAM_MAP:
        return NIZAM_MAP[nizam_name]
    upper = nizam_name.upper().strip()
    for key, val in NIZAM_MAP.items():
        if key.upper().strip() == upper:
            return val
    parts = upper.split()
    if len(parts) >= 2:
        for key, val in NIZAM_MAP.items():
            kp = key.upper().split()
            if len(kp) >= 2 and kp[0][:3] == parts[0][:3] and kp[1][:3] == parts[1][:3]:
                return val
    return None


def match_staff(conn, nizam_name):
    """staff.id for a NIZAM column header, or None. Writes nizam_name back on a
    map hit so the next import is an exact match."""
    row = conn.execute("SELECT id FROM staff WHERE nizam_name=?", (nizam_name,)).fetchone()
    if row:
        return row['id']
    row = conn.execute(
        "SELECT id FROM staff WHERE ulower(nizam_name) = ulower(?)"
        "   OR ulower(name) = ulower(?) OR ulower(full_name) = ulower(?)",
        (nizam_name, nizam_name, nizam_name)).fetchone()
    if row:
        conn.execute("UPDATE staff SET nizam_name=? WHERE id=? AND nizam_name IS NULL",
                     (nizam_name, row['id']))
        return row['id']
    mapped = _fuzzy_map(nizam_name)
    if mapped:
        row = conn.execute("SELECT id FROM staff WHERE name=?", (mapped,)).fetchone()
        if row:
            conn.execute("UPDATE staff SET nizam_name=? WHERE id=?", (nizam_name, row['id']))
            return row['id']
    return None


def read_nizam_sheet(filepath):
    """Parse the workbook into employees and project rows without touching
    the database, so a preview and the import read the same thing."""
    import openpyxl
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    employees = {}
    for col in range(4, ws.max_column + 1):
        v = ws.cell(2, col).value
        if v:
            employees[col] = str(v).replace('\n', ' ').strip()
    rows = []
    for r in range(4, ws.max_row + 1):
        name = ws.cell(r, 2).value
        total = parse_hhmm(ws.cell(r, 3).value)
        hours = {col: parse_hhmm(ws.cell(r, col).value) for col in employees}
        rows.append({'row': r, 'name': (str(name).strip() if name else ''),
                     'total': total, 'hours': hours})
    return {'employees': employees, 'rows': rows}


def import_nizam_file(filepath, period, create_projects=True):
    """Import one month of NIZAM hours. Returns a report dict.

    period is the YYYY-MM the sheet covers; every hour lands on that period.
    Re-importing the same month upserts — hours are replaced, phase tags and
    rate snapshots on the row are kept (an INSERT OR REPLACE would wipe them).
    """
    if not period or len(period) != 7 or period[4] != '-':
        raise ValueError('bad_period')
    sheet = read_nizam_sheet(filepath)
    employees = sheet['employees']
    file_name = os.path.basename(filepath)

    conn = get_db()
    try:
        staff_map, unmatched = {}, []
        for col, nizam_name in employees.items():
            sid = match_staff(conn, nizam_name)
            if sid:
                staff_map[col] = sid
            else:
                unmatched.append(nizam_name)

        projects_seen = set()
        projects_created = 0
        hour_entries = 0
        skipped = []
        unresolved = []
        for r in sheet['rows']:
            name = r['name']
            if not name:
                if r['total'] or any(r['hours'].values()):
                    skipped.append({'row': r['row'], 'reason': 'no_name', 'raw': ''})
                continue
            if any(kw in name.lower() for kw in NON_BILLABLE):
                skipped.append({'row': r['row'], 'reason': 'non_billable', 'raw': name})
                continue
            row_hours = sum(h for col, h in r['hours'].items() if col in staff_map)
            if r['total'] <= 0 and row_hours <= 0:
                skipped.append({'row': r['row'], 'reason': 'no_hours', 'raw': name})
                continue

            pid = resolve_entity_id(conn, 'project', name)
            if not pid:
                if not create_projects:
                    unresolved.append(name)
                    conn.execute(
                        "INSERT INTO unresolved_imports (entity_type, raw_name, field)"
                        " VALUES ('project', ?, 'nizam')", (name,))
                    continue
                cur = conn.execute("INSERT INTO projects (name, status) VALUES (?, 'active')",
                                   (name,))
                pid = cur.lastrowid
                projects_created += 1
                _write_audit(conn, 'create', 'projects', pid,
                             context=f'created by NIZAM import {file_name}')
            projects_seen.add(pid)

            for col, hrs in r['hours'].items():
                if hrs <= 0 or col not in staff_map:
                    continue
                conn.execute(
                    "INSERT INTO project_hours (project_id, staff_id, hours, period,"
                    " source, imported_at) VALUES (?,?,?,?,'nizam',?)"
                    " ON CONFLICT(project_id, staff_id, period)"
                    " DO UPDATE SET hours=excluded.hours, source='nizam',"
                    "   imported_at=excluded.imported_at",
                    (pid, staff_map[col], hrs, period, now_ts()))
                hour_entries += 1

        for sk in skipped:
            conn.execute(
                "INSERT INTO import_skip_log (file_name, sheet_name, row_num, reason, raw)"
                " VALUES (?,?,?,?,?)", (file_name, 'nizam', sk['row'], sk['reason'], sk['raw']))
        _write_audit(conn, 'import', 'project_hours', None,
                     context=f'NIZAM {file_name} period {period}: {hour_entries} rows,'
                             f' {len(projects_seen)} projects, {len(unmatched)} unmatched')
        conn.commit()
    finally:
        conn.close()

    return {
        'period': period,
        'projects': len(projects_seen),
        'projects_created': projects_created,
        'hour_entries': hour_entries,
        'skipped': skipped,
        'employees_mapped': len(staff_map),
        'unmatched_names': unmatched,
        'unresolved_projects': unresolved,
    }
