# -*- coding: utf-8 -*-
"""End-to-end smoke test — every page renders, and a full document flow works
through the HTTP layer rather than by calling the models directly.

Uses Flask's test client, so no server needs to be running.

    python test_app.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

DB_FILE = os.path.join(tempfile.gettempdir(), 'mizan5_test_app.db')
os.environ['MIZAN5_DB'] = DB_FILE
os.environ['MIZAN_ADMIN_PASSWORD'] = 'testpass'
os.environ['MIZAN_SECRET_KEY'] = 'test-secret'
if os.path.exists(DB_FILE):
    os.remove(DB_FILE)

from models.base import init_db, get_db                              # noqa: E402
from models.ledger import get_trial_balance, account_balance         # noqa: E402
from app import create_app                                           # noqa: E402

PASS, FAIL = 0, 0


def check(label, condition, detail=''):
    global PASS, FAIL
    if condition:
        PASS += 1
        print(f'  [ok]   {label}')
    else:
        FAIL += 1
        print(f'  [FAIL] {label}' + (f'  — {detail}' if detail else ''))


init_db()
conn = get_db()
CLIENT = conn.execute(
    "INSERT INTO counterparties (name, inn, counterparty_type)"
    " VALUES ('Alfa MChJ','301234567','client')").lastrowid
VENDOR = conn.execute(
    "INSERT INTO counterparties (name, counterparty_type) VALUES ('Beta MChJ','vendor')"
).lastrowid
PROJECT = conn.execute(
    "INSERT INTO projects (name, counterparty_id, is_billable, estimated_total_hours)"
    " VALUES ('Biznes markaz', ?, 1, 800)", (CLIENT,)).lastrowid
SID = conn.execute(
    "INSERT INTO staff (name, role, department, staff_type)"
    " VALUES ('Aziz','Arxitektor','Arxitektura','production')").lastrowid
conn.execute("INSERT INTO salary_history (staff_id, base_salary, premium, start_date)"
             " VALUES (?, 9000000, 1000000, '2026-01-01')", (SID,))
conn.execute("INSERT INTO project_hours (project_id, staff_id, hours, period)"
             " VALUES (?,?,?,?)", (PROJECT, SID, 120, '2026-06'))
LOAN = conn.execute(
    "INSERT INTO loans (loan_type, counterparty_id, total_amount, issue_date)"
    " VALUES ('olgan', ?, 50000000, '2026-01-10')", (VENDOR,)).lastrowid
# Equipment so the depreciation card on /periods has real rows to render, and
# one dateless asset so the "skipped" block renders too.
conn.execute(
    "INSERT INTO equipment (name, kind, staff_id, quantity, price,"
    " lifespan_months, purchase_date) VALUES ('Laptop','personal',?,1,15000000,36,"
    " '2026-01-15')", (SID,))
conn.execute(
    "INSERT INTO equipment (name, kind, quantity, price, lifespan_months,"
    " purchase_date) VALUES ('Server','general',1,30000000,60,'2026-01-10')")
conn.execute(
    "INSERT INTO equipment (name, kind, quantity, price, lifespan_months)"
    " VALUES ('Monitor','general',1,4000000,24)")
conn.commit()
conn.close()

app = create_app()
app.config['WTF_CSRF_ENABLED'] = False
# Uploads (the NIZAM import test) go to the temp dir, never the project folder.
app.config['UPLOAD_FOLDER'] = os.path.join(tempfile.gettempdir(), 'mizan5_test_uploads')
client = app.test_client()

print('\n=== Authentication ===')
r = client.get('/')
check('anonymous request is redirected to login',
      r.status_code in (302, 401), str(r.status_code))
r = client.post('/login', data={'username': 'admin', 'password': 'wrong'})
check('a wrong password does not log in', b'MIZAN' in r.data and r.status_code == 200)
r = client.post('/login', data={'username': 'admin', 'password': 'testpass'},
                follow_redirects=True)
check('correct credentials log in', r.status_code == 200 and b'MIZAN' in r.data)

print('\n=== Every page renders ===')
PAGES = [
    ('/', 'dashboard'),
    ('/documents/sales_invoice', 'sales invoices'),
    ('/documents/purchase_invoice', 'purchase invoices'),
    ('/documents/cash_in', 'money in'),
    ('/documents/cash_out', 'money out'),
    ('/documents/cash_out?direction=internal', 'money out filtered by direction'),
    ('/documents/sales_invoice?direction=external', 'invoices filtered by direction'),
    ('/documents/manual', 'manual entry'),
    ('/documents/opening', 'opening balances'),
    ('/documents/dividend', 'dividends'),
    ('/journal', 'journal'),
    ('/trial-balance', 'trial balance'),
    ('/accounts', 'chart of accounts'),
    ('/counterparties', 'counterparties'),
    ('/bank-accounts', 'bank accounts'),
    ('/staff', 'staff'),
    ('/staff/hours', 'timesheets'),
    ('/staff/equipment', 'equipment'),
    ('/rates', 'rates'),
    ('/rates/overhead', 'overhead reconciliation'),
    ('/projects', 'projects'),
    (f'/projects/{PROJECT}', 'project detail'),
    ('/budget', 'budget'),
    ('/payroll', 'payroll'),
    ('/loans', 'loans'),
    ('/reports/pnl', 'profit and loss'),
    ('/reports/balance-sheet', 'balance sheet'),
    ('/reports/cashflow', 'cash flow'),
    ('/reports/vat', 'VAT report'),
    ('/reports/aging', 'AR aging'),
    ('/reports/aging?kind=ap', 'AP aging'),
    ('/periods', 'periods'),
    ('/settings', 'settings'),
    # v4 feature port
    ('/kpi', 'staff KPI'),
    ('/pricing', 'pricing / scratch quote'),
    ('/projects?sort=name&status=active&q=biz', 'projects filtered and sorted'),
    ('/documents/sales_invoice?sort=total&dir=asc&page=1&per_page=25', 'invoices sorted and paged'),
    (f'/documents/sales_invoice?project_id={PROJECT}&responsible_id={SID}', 'invoices by project and responsible'),
    ('/staff?department=Arxitektura&staff_type=production&active=1&q=az', 'staff filtered'),
    ('/staff/equipment?kind=personal&asset_class=computer&license_type=named', 'equipment filtered'),
    ('/staff/hours?period=2026-06', 'timesheets for a period'),
]
for url, label in PAGES:
    r = client.get(url)
    ok = r.status_code == 200
    detail = ''
    if not ok:
        detail = f'HTTP {r.status_code}'
        body = r.data.decode('utf-8', 'replace')
        for marker in ('UndefinedError', 'TemplateSyntaxError', 'jinja2', 'Error'):
            idx = body.find(marker)
            if idx > -1:
                detail += ' — ' + body[idx:idx + 180].replace('\n', ' ')
                break
    check(f'{label} ({url})', ok, detail)

print('\n=== Document flow through HTTP ===')
r = client.post('/documents/sales_invoice/save', data={
    'date': '2026-06-10', 'counterparty_id': CLIENT, 'project_id': PROJECT,
    'due_date': '2026-07-10', 'description': 'Loyiha ishlari', 'currency': 'UZS',
    'line_description': ['Arxitektura loyihasi'], 'line_quantity': ['1'],
    'line_unit_price': ['20000000'], 'line_amount': ['20000000'],
    'line_vat_rate': ['12'], 'line_vat_amount': ['2400000'],
    'post_now': '1',
}, follow_redirects=True)
check('sales invoice posts through the form', r.status_code == 200)
check('AR was debited with the gross total',
      abs(account_balance('4010') - 22400000) < 1, str(account_balance('4010')))
check('revenue was credited net of VAT',
      abs(account_balance('9030') - 20000000) < 1, str(account_balance('9030')))
check('output VAT was recorded',
      abs(account_balance('6410.1') - 2400000) < 1)

conn = get_db()
inv_id = conn.execute("SELECT id FROM documents WHERE doc_type='sales_invoice'"
                      " ORDER BY id DESC LIMIT 1").fetchone()['id']
conn.close()

r = client.post('/documents/cash_in/save', data={
    'date': '2026-06-25', 'counterparty_id': CLIENT, 'payment_method': 'bank',
    'total': '10000000', 'currency': 'UZS',
    'alloc_invoice_doc_id': [str(inv_id)], 'alloc_amount': ['10000000'],
    'post_now': '1',
}, follow_redirects=True)
check('payment posts and allocates', r.status_code == 200)
check('cash increased', abs(account_balance('5110') - 10000000) < 1,
      str(account_balance('5110')))
check('AR reduced by the allocation',
      abs(account_balance('4010') - 12400000) < 1, str(account_balance('4010')))

r = client.get('/reports/aging')
check('the unpaid balance shows in aging', b'12,400,000' in r.data)

print('\n=== Payroll through HTTP ===')
r = client.post('/payroll/save', data={
    'period': '2026-06', 'row_staff_id': [str(SID)], 'row_gross': ['10000000'],
    'row_pit': ['1200000'], 'row_social': ['1200000'], 'post_now': '1',
}, follow_redirects=True)
check('payroll accrues through the form', r.status_code == 200)
check('net pay is owed to staff', abs(account_balance('6710') - 8800000) < 1,
      str(account_balance('6710')))
check('production labor hit 2010',
      abs(account_balance('2010') - 11200000) < 1, str(account_balance('2010')))

r = client.post('/payroll/remit', data={'period': '2026-06', 'date': '2026-07-05',
                                        'payment_method': 'bank'},
                follow_redirects=True)
check('remittance clears the payroll liabilities',
      abs(account_balance('6710')) < 1 and abs(account_balance('6420.1')) < 1
      and abs(account_balance('6520')) < 1,
      f"{account_balance('6710')}/{account_balance('6420.1')}/{account_balance('6520')}")

print('\n=== Dividend declaration through HTTP ===')
r = client.get('/documents/dividend')
check('dividend page is linked from the sidebar', b'href="/documents/dividend"' in r.data)
check('dividend modal offers an amount input',
      b'name="total"' in r.data, 'no total field rendered')
r = client.post('/documents/dividend/save', data={
    'date': '2026-06-30', 'counterparty_id': CLIENT, 'total': '5000000',
    'currency': 'UZS', 'description': 'Dividend 2026 H1', 'post_now': '1',
}, follow_redirects=True)
check('dividend declaration posts through the form',
      r.status_code == 200 and b'class="alert alert-success"' in r.data)
# Both accounts are credit-normal, so the Dr 8710 shows as a negative balance
# and the Cr 6610 as a positive one.
check('retained earnings debited', abs(account_balance('8710') + 5000000) < 1,
      str(account_balance('8710')))
check('dividends payable credited', abs(account_balance('6610') - 5000000) < 1,
      str(account_balance('6610')))

print('\n=== Manual entry balance guard ===')
r = client.post('/documents/manual/save', data={
    'date': '2026-06-30', 'description': 'Test',
    'line_account_id': ['1', '2'], 'line_debit': ['100', ''],
    'line_credit': ['', '40'], 'post_now': '1',
}, follow_redirects=True)
check('an unbalanced manual entry is rejected with a message',
      b'alert-error' in r.data or b'ledger_unbalanced' in r.data
      or 'balanslashmagan'.encode() in r.data, 'no error surfaced')

print('\n=== Milestones and pricing through HTTP ===')
r = client.post(f'/projects/{PROJECT}/schedule', data={}, follow_redirects=True)
check('standard schedule is created', r.status_code == 200)
conn = get_db()
phase = conn.execute("SELECT id FROM project_phases WHERE project_id=?"
                     " ORDER BY sort_order LIMIT 1", (PROJECT,)).fetchone()
conn.close()
check('milestones exist', phase is not None)

if phase:
    r = client.post(f'/projects/{PROJECT}/milestone', data={
        'id': str(phase['id']), 'name': 'Eskiz loyiha', 'work_type': 'eskiz',
        'start_date': '2026-03-01', 'end_date': '2026-05-31',
        'planned_outsourcing': '2000000', 'status': 'in_progress',
        'has_staff_grid': '1',
        'ms_staff_id': [str(SID)], 'ms_est_hours': ['150'],
    }, follow_redirects=True)
    check('milestone saves with its staff hours', r.status_code == 200)
    conn = get_db()
    row = conn.execute("SELECT planned_hours, planned_cost FROM project_phases"
                       " WHERE id=?", (phase['id'],)).fetchone()
    conn.close()
    check('planned hours derived from the staff grid', row['planned_hours'] == 150,
          str(row['planned_hours']))
    check('planned cost derived from the rate', row['planned_cost'] > 0,
          str(row['planned_cost']))

r = client.post(f'/projects/{PROJECT}/prices', data={}, follow_redirects=True)
check('suggested prices apply', r.status_code == 200)
r = client.post(f'/projects/{PROJECT}/freeze', data={}, follow_redirects=True)
check('plan freezes', r.status_code == 200)
r = client.get('/budget')
check('budget page shows the frozen project', b'Biznes markaz' in r.data)

print('\n=== Counterparty and staff modals ===')
r = client.post('/counterparties/save', data={
    'name': 'Gamma Qurilish', 'inn': '309999999', 'counterparty_type': 'client',
    'is_active': '1'}, follow_redirects=True)
check('counterparty saves', b'Gamma Qurilish' in r.data)
r = client.post('/staff/save', data={
    'name': 'Bekzod', 'role': 'BIM muhandis', 'department': 'Arxitektura',
    'staff_type': 'production', 'is_active': '1'}, follow_redirects=True)
check('staff member saves', b'Bekzod' in r.data)
r = client.post('/staff/salary', data={
    'staff_id': str(SID), 'start_date': '2026-07-01',
    'base_salary': '11000000', 'premium': '0'}, follow_redirects=True)
check('salary change is recorded', r.status_code == 200)
conn = get_db()
closed = conn.execute("SELECT COUNT(*) AS n FROM salary_history"
                      " WHERE staff_id=? AND end_date IS NOT NULL", (SID,)).fetchone()['n']
conn.close()
check('the previous salary row was closed', closed == 1, str(closed))

print('\n=== Period close through HTTP ===')
print('\n=== Depreciation through HTTP ===')
r = client.get('/periods?dep_period=2026-05')
check('the depreciation card renders its schedule',
      r.status_code == 200 and 'Laptop'.encode() in r.data)
check('a dateless asset is reported as skipped, not dropped',
      'Monitor'.encode() in r.data)
r = client.post('/periods/depreciate', data={'period': '2026-05'},
                follow_redirects=True)
check('depreciation posts through the form', r.status_code == 200)
from models.ledger import account_balance as _bal                     # noqa: E402
# The credit lands on the contra account of the asset's CLASS (0250 for
# computers), not on the 0200 parent, so sum the 02xx family.
_accum = sum(_bal(c) for c in ('0200', '0220.1', '0230', '0240', '0250',
                               '0260', '0290'))
check('accumulated depreciation is now on the books', _accum > 0, str(_accum))
check('it is credited per asset class, not to the 0200 parent',
      _bal('0250') > 0 and abs(_bal('0200')) < 1,
      f"0250={_bal('0250')} 0200={_bal('0200')}")
check('the expense landed on 9420.1, not 9420',
      _bal('9420.1') > 0 and abs(_bal('9420.1') - _accum) < 1,
      f"9420.1={_bal('9420.1')} accum={_accum}")
r = client.post('/periods/depreciate', data={'period': '2026-05'},
                follow_redirects=True)
check('posting the same period twice is refused politely',
      b'alert-info' in r.data or b'alert-success' in r.data)

r = client.post('/periods/close', data={'period': '2026-06', 'status': 'hard_closed'},
                follow_redirects=True)
check('period closes', r.status_code == 200)
r = client.post('/documents/sales_invoice/save', data={
    'date': '2026-06-15', 'counterparty_id': CLIENT,
    'line_description': ['X'], 'line_amount': ['1000'], 'line_vat_rate': ['0'],
    'post_now': '1'}, follow_redirects=True)
check('posting into the closed period is refused with a message',
      b'alert-error' in r.data, 'no error surfaced')
r = client.post('/periods/reopen', data={'period': '2026-06'}, follow_redirects=True)
check('period reopens', r.status_code == 200)

print('\n=== Language switching ===')
for code, marker in (('ru', 'Бухгалтерия'), ('en', 'Accounting'), ('uz', 'Buxgalteriya')):
    client.get(f'/lang/{code}')
    r = client.get('/')
    check(f'interface switches to {code}', marker.encode('utf-8') in r.data)

# Uzbek text is full of apostrophes (o'chirish, bo'lmaydi). Dropped into a
# single-quoted JS literal inside an onsubmit or <script>, one apostrophe ends
# the string, the handler throws, and the form either submits unconfirmed or —
# when the handler was also setting the action — POSTs to a GET-only URL (405).
import re
client.get('/lang/uz')
# A bank account under 5110, registered after the cash_in above was posted, so
# the page shows an Unassigned row with its Assign form.
conn = get_db()
ACC_5110 = conn.execute("SELECT id FROM accounts WHERE code='5110'").fetchone()['id']
conn.close()
r = client.post('/bank-accounts/save', data={
    'account_id': ACC_5110, 'name': 'Asosiy hisob', 'account_number': '20208000900000000001',
    'bank_name': 'Kapitalbank', 'currency': 'UZS', 'is_default': '1', 'is_active': '1',
}, follow_redirects=True)
check('bank account registers through the form', r.status_code == 200)
body = client.get('/bank-accounts').data.decode('utf-8')
check('unassigned 5110 lines offer an Assign form', '/assign-unassigned' in body)

# A quoted JS literal with a word-internal apostrophe (o'ch, bo'l) inside it.
bad_js = re.compile(r"""(confirm\(|textContent\s*=\s*)'[^'\n]*[A-Za-z]'[a-z]""")
for url in ('/documents/cash_out', f'/documents/{inv_id}', '/periods',
            f'/projects/{PROJECT}', '/bank-accounts', '/settings'):
    body = client.get(url).data.decode('utf-8')
    m = bad_js.search(body)
    check(f'no apostrophe breaks inline JS on {url} in Uzbek', m is None,
          m.group(0)[:100] if m else '')

# The Assign form must post to the assign route, not to the GET-only list.
conn = get_db()
BANK = conn.execute("SELECT id FROM bank_accounts WHERE account_id=?", (ACC_5110,)).fetchone()['id']
conn.close()
r = client.post(f'/bank-accounts/{BANK}/assign-unassigned', follow_redirects=True)
check('assign-unassigned posts through HTTP',
      r.status_code == 200 and b'class="alert alert-success"' in r.data)
body = client.get('/bank-accounts').data.decode('utf-8')
check('nothing is left unassigned on 5110', '/assign-unassigned' not in body)

print('\n=== v4 feature port through HTTP ===')
# Scratch quote API — stateless, then saved into a project as milestones.
quote_body = {'rows': [{'name': 'Eskiz', 'work_type': 'eskiz', 'start': '2027-01', 'end': '2027-02',
                        'outsourcing': 500000, 'staff': [{'staff_id': SID, 'hours': 40}]}],
              'risk': {'deadline_months': 6, 'client_type': 'new', 'complexity': 'medium',
                       'currency': 'UZS'}}
r = client.post('/api/pricing/quote', json=quote_body)
check('scratch quote returns a priced schedule',
      r.status_code == 200 and r.get_json()['quote']['target_contract'] > 0)
conn = get_db()
QP = conn.execute("INSERT INTO projects (name, is_billable) VALUES ('Quote target', 1)").lastrowid
conn.commit()
conn.close()
r = client.post('/api/pricing/quote/save', json={**quote_body, 'project_id': QP})
check('quote saves into a project', r.status_code == 200 and r.get_json()['created'] == 1, r.data[:120])
r = client.post('/api/pricing/quote/save', json={**quote_body, 'project_id': QP})
check('saving twice without overwrite is refused', r.status_code == 400 and r.get_json()['error'] == 'exists')
r = client.get(f'/projects/{QP}')
check('project page renders the quoted milestone', r.status_code == 200 and b'Eskiz' in r.data)

# Unassigned hours → milestone from the project page.
conn = get_db()
PH = conn.execute("SELECT id FROM project_phases WHERE project_id=? LIMIT 1", (PROJECT,)).fetchone()
conn.close()
if PH:
    r = client.post(f'/projects/{PROJECT}/assign-hours',
                    data={'period': '2026-06', 'phase_id': PH['id']}, follow_redirects=True)
    conn = get_db()
    tagged = conn.execute("SELECT phase_id FROM project_hours WHERE project_id=? AND period='2026-06'",
                          (PROJECT,)).fetchone()['phase_id']
    conn.close()
    check('assign-hours tags the month', r.status_code == 200 and tagged == PH['id'])
    r = client.post(f'/projects/{PROJECT}/assign-hours',
                    data={'period': '2026-06', 'phase_id': ''}, follow_redirects=True)
    check('blank phase unassigns again', r.status_code == 200)

# NIZAM upload through the form.
import io as _io                                                      # noqa: E402
import openpyxl                                                       # noqa: E402
wb = openpyxl.Workbook()
ws = wb.active
ws.cell(2, 4, 'Aziz')
ws.cell(2, 5, 'GHOST EMPLOYEE')
ws.cell(4, 2, 'Biznes markaz'); ws.cell(4, 3, '8:00'); ws.cell(4, 4, '8:00')
ws.cell(5, 2, 'HR'); ws.cell(5, 3, '2:00'); ws.cell(5, 4, '2:00')
buf = _io.BytesIO()
wb.save(buf)
buf.seek(0)
r = client.post('/staff/hours/import',
                data={'file': (buf, 'table.xlsx'), 'period': '2026-07', 'create_projects': '1'},
                content_type='multipart/form-data', follow_redirects=True)
check('NIZAM upload imports and reports', r.status_code == 200
      and 'GHOST EMPLOYEE'.encode() in r.data and b'2026-07' in r.data)
conn = get_db()
imported = conn.execute("SELECT hours, source FROM project_hours WHERE project_id=? AND staff_id=?"
                        " AND period='2026-07'", (PROJECT, SID)).fetchone()
conn.close()
check('imported hours landed on the period', imported and imported['hours'] == 8
      and imported['source'] == 'nizam')
r = client.post('/staff/hours/import', data={'period': '2026-07'},
                content_type='multipart/form-data', follow_redirects=True)
check('an upload without a file is rejected politely', r.status_code == 200)

# Periods: create, delete, and the hard-close lock.
r = client.post('/periods/create', data={'code': '2032-01', 'notes': 'ahead'}, follow_redirects=True)
check('period created from the form', r.status_code == 200 and b'2032-01' in r.data)
r = client.post('/periods/notes', data={'period': '2032-01', 'notes': 'edited'}, follow_redirects=True)
check('period notes saved', r.status_code == 200 and b'edited' in r.data)
r = client.post('/periods/delete', data={'period': '2032-01'}, follow_redirects=True)
check('empty period deleted', r.status_code == 200 and b'2032-01' not in r.data)
from models.base import get_period_status                             # noqa: E402
from models.reports import list_periods                               # noqa: E402
hard = [p for p in list_periods() if p['status'] == 'hard_closed']
if hard:
    r = client.post('/periods/reopen', data={'period': hard[0]['code']}, follow_redirects=True)
    check('hard-closed period stays closed through the UI',
          r.status_code == 200 and get_period_status(hard[0]['code']) == 'hard_closed')

# Loans: repayment through the form and auto-close. 5110 is subdivided by now
# (a bank account was registered above), so bank money must name its account.
r = client.post('/loans/save', data={'loan_type': 'olgan', 'counterparty_id': VENDOR,
                                     'total_amount': '3000000', 'currency': 'UZS',
                                     'issue_date': '2026-08-01', 'term': 'short',
                                     'status': 'ochiq', 'post_now': '1',
                                     'payment_method': 'bank', 'bank_account_id': BANK},
                follow_redirects=True)
check('loan booked through the form', r.status_code == 200)
conn = get_db()
L2 = conn.execute("SELECT id FROM loans WHERE total_amount=3000000 ORDER BY id DESC LIMIT 1").fetchone()['id']
conn.close()
r = client.post(f'/loans/{L2}/payment', data={'principal': '3000000', 'interest': '100000',
                                              'date': '2026-08-20', 'payment_method': 'bank',
                                              'bank_account_id': BANK},
                follow_redirects=True)
conn = get_db()
st = conn.execute("SELECT status FROM loans WHERE id=?", (L2,)).fetchone()['status']
conn.close()
check('a full repayment closes the loan', r.status_code == 200 and st == 'yopilgan', st)
check('loans page shows the payment history', b'hist' + str(L2).encode() in client.get('/loans').data)

# Documents list: paging and sorting survive the round trip.
r = client.get('/documents/sales_invoice?sort=total&dir=asc&per_page=25&page=1')
check('sorted, paged invoice list renders with a pager',
      r.status_code == 200 and b'per_page' in r.data)
r = client.get('/documents/sales_invoice?sort=nonsense&page=999')
check('bad sort and page are clamped, not 500', r.status_code == 200)

# Settings: nobody can lock themselves out; the CBU endpoint validates input.
conn = get_db()
ME = conn.execute("SELECT id FROM users WHERE username='admin'").fetchone()['id']
conn.close()
r = client.post('/settings/user', data={'id': ME, 'username': 'admin', 'role': 'viewer',
                                        'is_active': '1'}, follow_redirects=True)
conn = get_db()
role = conn.execute("SELECT role, is_active FROM users WHERE id=?", (ME,)).fetchone()
conn.close()
check('admin cannot demote or deactivate themselves',
      role['role'] == 'admin' and role['is_active'] == 1)
r = client.get('/settings/fetch-rate')
check('CBU fetch without a date is a 400', r.status_code == 400)

# KPI page renders with data.
r = client.get('/kpi')
check('KPI page lists the production employee', r.status_code == 200 and b'Aziz' in r.data)

print('\n=== Books remain intact ===')
check('trial balance is balanced', get_trial_balance()['is_balanced'])
from models.ledger import verify_all_entries                          # noqa: E402
check('no unbalanced entry exists', not verify_all_entries())
from models.reports import get_balance_sheet                          # noqa: E402
check('balance sheet balances', get_balance_sheet()['is_balanced'])

print('\n=== Logout ===')
r = client.get('/logout', follow_redirects=True)
check('logout returns to the login page', b'MIZAN' in r.data)
r = client.get('/rates')
check('protected pages are closed after logout', r.status_code in (302, 401))

print(f'\n{"=" * 46}\n  passed: {PASS}   failed: {FAIL}\n{"=" * 46}')
sys.exit(1 if FAIL else 0)
