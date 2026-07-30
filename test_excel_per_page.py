"""Regression test: Excel ?sheet= parametri mos sahifani aktiv qiladi.

Run: python test_excel_per_page.py
Bug: /staff sahifasidan Excel tugmasi bosilganda fayl Dashboard tab-da ochilar edi
     (chunki butun multi-sheet workbook qaytarilardi va Dashboard birinchi).
Fix: ?sheet=<name> query param qabul qilinadi, mos sheet wb.active qilib o'rnatiladi.
"""
import sys, os, urllib.request, tempfile
sys.path.insert(0, os.path.dirname(__file__))

# Need running server — caller must start app on PORT 5099 first
BASE = 'http://127.0.0.1:5099'

EXPECTED = {
    'staff': '03-Maosh',
    'hourly': '07-Soat narxi',
    'budget': '08-Byudjet',
    'kpi': '10-KPI',
    'dashboard': '09-Dashboard',
}


def fetch_export(query=''):
    url = f'{BASE}/api/export/xlsx{query}'
    fd, path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    urllib.request.urlretrieve(url, path)
    return path


def get_active_sheet(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    title = wb.active.title
    wb.close()
    return title


def test_default_no_param_active_dashboard():
    """No ?sheet param → Dashboard birinchi (eski xulq saqlangan)."""
    path = fetch_export()
    active = get_active_sheet(path)
    os.unlink(path)
    assert active == '09-Dashboard', f"default: kutilgan 09-Dashboard, olindi {active}"
    print("  [OK] test_default_no_param_active_dashboard")


def test_per_page_sheets():
    """Har bir ?sheet=<name> mos sahifani aktiv qiladi."""
    for param, expected_title in EXPECTED.items():
        path = fetch_export(f'?sheet={param}')
        active = get_active_sheet(path)
        os.unlink(path)
        assert active == expected_title, \
            f"sheet={param}: kutilgan {expected_title}, olindi {active}"
    print(f"  [OK] test_per_page_sheets ({len(EXPECTED)} ta sahifa)")


def test_invalid_sheet_param_falls_back():
    """Noto'g'ri ?sheet=xyz → default Dashboard aktiv (graceful fallback)."""
    path = fetch_export('?sheet=invalid_xyz')
    active = get_active_sheet(path)
    os.unlink(path)
    assert active == '09-Dashboard', \
        f"invalid sheet: kutilgan default Dashboard, olindi {active}"
    print("  [OK] test_invalid_sheet_param_falls_back")


def test_all_sheets_still_present():
    """?sheet=staff bo'lsa ham, faylda hamma 5 sahifa bor (faqat aktiv farqi)."""
    import openpyxl
    path = fetch_export('?sheet=staff')
    wb = openpyxl.load_workbook(path, read_only=True)
    assert len(wb.sheetnames) == 5, f"sahifalar soni: kutilgan 5, olindi {len(wb.sheetnames)}"
    for expected in EXPECTED.values():
        assert expected in wb.sheetnames, f"sahifa yo'q: {expected}"
    wb.close()
    os.unlink(path)
    print("  [OK] test_all_sheets_still_present")


if __name__ == '__main__':
    print("Running Excel per-page regression tests...")
    print("(App :5099-portda ishga tushgan bo'lishi kerak)")
    try:
        test_default_no_param_active_dashboard()
        test_per_page_sheets()
        test_invalid_sheet_param_falls_back()
        test_all_sheets_still_present()
        print("\nAll tests passed.")
    except AssertionError as e:
        print(f"\nFAIL: {e}")
        sys.exit(1)
    except Exception as e:
        print(f"\nERROR: {type(e).__name__}: {e}")
        raise
